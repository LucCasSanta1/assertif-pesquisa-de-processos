from __future__ import annotations

import logging
import os
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from src.logic.archive_classifier import classificar_arquivamento
from src.logic.court_classifier import classificar_esfera_e_tribunal, formatar_numero_cnj
from src.models.schemas import (
    EsferaJustica,
    MovimentoProcessual,
    ProcessoJudicial,
    StatusArquivamento,
)

logger = logging.getLogger(__name__)

# Diretórios onde planilhas base (.xlsx) podem ser lidas automaticamente
BASE_DIR = Path(__file__).resolve().parent.parent.parent
DATA_DIR = BASE_DIR / "data"
DOWNLOADS_SUDATI = Path(r"C:\Users\asser\Downloads\Grupo Sudati 2.0.xlsx")

# Cache em memória para leitura instantânea após o primeiro carregamento
_CACHE_PLANILHA: Dict[str, List[ProcessoJudicial]] = {}
_ARQUIVOS_CARREGADOS: set[str] = set()


def _limpar_digitos(val: object) -> str:
    if val is None:
        return ""
    return re.sub(r"\D", "", str(val))


def _parse_data_planilha(val: object) -> Optional[datetime]:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val
    s = str(val).strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%d/%m/%Y", "%d/%m/%Y %H:%M:%S"):
        try:
            return datetime.strptime(s[:19], fmt)
        except ValueError:
            continue
    return None


def _mapear_status_planilha(
    status_raw: str,
    motivo_raw: str,
    data_arquivo: Optional[datetime],
) -> Tuple[bool, StatusArquivamento, Optional[int], Optional[str], List[MovimentoProcessual]]:
    s = (status_raw or "").upper().strip()
    motivo = (motivo_raw or "").strip()
    dt = data_arquivo or datetime.now()
    movs: List[MovimentoProcessual] = []

    if "DEFINITIVO" in s or "BAIXA" in s:
        desc = f"{status_raw.strip()}" + (f" ({motivo})" if motivo else "")
        movs.append(MovimentoProcessual(codigo=246, nome=desc, data_hora=dt))
        return True, StatusArquivamento.ARQUIVADO_DEFINITIVO, 246, desc, movs

    if "PROVISORIO" in s or "PROVISÓRIO" in s:
        desc = f"{status_raw.strip()}" + (f" ({motivo})" if motivo else "")
        movs.append(MovimentoProcessual(codigo=245, nome=desc, data_hora=dt))
        return True, StatusArquivamento.ARQUIVADO_PROVISORIO, 245, desc, movs

    if "ARQUIVAMENTO" in s or "ARQUIVADO" in s:
        desc = f"{status_raw.strip()}" + (f" ({motivo})" if motivo else "")
        movs.append(MovimentoProcessual(codigo=861, nome=desc, data_hora=dt))
        return True, StatusArquivamento.ARQUIVADO_DEFINITIVO, 861, desc, movs

    if s:
        return False, StatusArquivamento.ATIVO, None, status_raw.strip(), movs

    return False, StatusArquivamento.INDETERMINADO, None, None, movs


def carregar_bytes_planilha_para_cache(conteudo_xlsx: bytes, nome_origem: str = "Planilha Importada") -> int:
    """
    Lê uma planilha Excel (.xlsx) enviada pelo usuário (ex: Grupo Sudati 2.0.xlsx)
    e indexa seus processos por CNPJ (14 dígitos) em memória.
    Retorna a quantidade total de processos indexados.
    """
    wb = load_workbook(BytesIO(conteudo_xlsx), read_only=True, data_only=True)
    total = 0
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c or "").upper().strip() for c in rows[0]]

        idx_cnpj = next((i for i, h in enumerate(header) if "CNPJ" in h), None)
        idx_proc = next((i for i, h in enumerate(header) if "PROCESSO" in h), None)
        if idx_cnpj is None or idx_proc is None:
            continue

        idx_razao = next((i for i, h in enumerate(header) if "RAZ" in h and "SOCIAL" in h), None)
        idx_status = next((i for i, h in enumerate(header) if h == "STATUS" or "STATUS" in h), None)
        idx_dt_arq = next((i for i, h in enumerate(header) if "DATA ARQUIVO" in h or "DATA ARQUIVAMENTO" in h), None)
        idx_motivo = next((i for i, h in enumerate(header) if "MOTIVO" in h), None)
        idx_trib = next((i for i, h in enumerate(header) if "TRIBUNAL" in h), None)
        idx_uf = next((i for i, h in enumerate(header) if h == "UF"), None)
        idx_vara = next((i for i, h in enumerate(header) if "VARA" in h or "ORGAO" in h or "ÓRGÃO" in h), None)
        idx_parte = next((i for i, h in enumerate(header) if "PARTE" in h), None)

        for r in rows[1:]:
            if not r or idx_cnpj >= len(r) or idx_proc >= len(r):
                continue
            cnpj_14 = _limpar_digitos(r[idx_cnpj])
            if len(cnpj_14) != 14:
                continue
            num_20 = _limpar_digitos(r[idx_proc])
            if len(num_20) != 20:
                continue

            trib_raw = str(r[idx_trib] or "").strip().upper().replace("-", "") if idx_trib is not None and idx_trib < len(r) else None
            esfera, sigla_padrao, uf_derivada = classificar_esfera_e_tribunal(num_20, None)
            if trib_raw and not trib_raw.startswith("JF"):
                sigla_padrao = trib_raw

            uf_val = (
                str(r[idx_uf]).strip().upper()
                if idx_uf is not None and idx_uf < len(r) and r[idx_uf]
                else uf_derivada
            )
            vara_val = (
                str(r[idx_vara]).strip()
                if idx_vara is not None and idx_vara < len(r) and r[idx_vara]
                else None
            )
            parte_val = (
                str(r[idx_parte]).strip()
                if idx_parte is not None and idx_parte < len(r) and r[idx_parte]
                else None
            )
            razao_val = (
                str(r[idx_razao]).strip()
                if idx_razao is not None and idx_razao < len(r) and r[idx_razao]
                else None
            )
            status_raw = (
                str(r[idx_status]).strip()
                if idx_status is not None and idx_status < len(r) and r[idx_status]
                else ""
            )
            motivo_raw = (
                str(r[idx_motivo]).strip()
                if idx_motivo is not None and idx_motivo < len(r) and r[idx_motivo]
                else ""
            )
            dt_arq = (
                _parse_data_planilha(r[idx_dt_arq])
                if idx_dt_arq is not None and idx_dt_arq < len(r)
                else None
            )

            arq, status_det, cod_mov, desc_mov, movs = _mapear_status_planilha(status_raw, motivo_raw, dt_arq)

            # Data de ajuizamento estimada pelo ano do CNJ se não houver
            ano_cnj = int(num_20[9:13]) if num_20[9:13].isdigit() else None
            dt_ajuiz = datetime(ano_cnj, 1, 1) if ano_cnj and 1990 <= ano_cnj <= 2030 else None

            partes_str = " | ".join(filter(None, [razao_val, f"Parte Contrária: {parte_val}" if parte_val else None])) or None

            proc = ProcessoJudicial(
                numero_processo=num_20,
                numero_formatado=formatar_numero_cnj(num_20),
                tribunal=sigla_padrao,
                esfera=esfera,
                orgao_julgador=vara_val,
                uf=uf_val,
                data_ajuizamento=dt_ajuiz,
                data_ultima_movimentacao=dt_arq,
                ultimo_movimento_relevante=desc_mov,
                codigo_movimento_arquivamento=cod_mov,
                data_arquivamento=dt_arq if arq else None,
                arquivado=arq,
                status_detalhado=status_det,
                polo_empresa="Reclamado / Parte",
                partes_resumo=partes_str,
                fonte_descoberta=f"Base Acervo ({nome_origem})",
                movimentos=movs,
            )
            _CACHE_PLANILHA.setdefault(cnpj_14, []).append(proc)
            total += 1

    wb.close()
    return total


def _garantir_bases_locais_carregadas() -> None:
    candidatos: List[Path] = []
    if DATA_DIR.exists():
        candidatos.extend(DATA_DIR.glob("*.xlsx"))
    if DOWNLOADS_SUDATI.exists():
        candidatos.append(DOWNLOADS_SUDATI)

    for caminho in candidatos:
        chave = str(caminho.resolve())
        if chave in _ARQUIVOS_CARREGADOS or caminho.name.startswith("~$"):
            continue
        try:
            conteudo = caminho.read_bytes()
            carregar_bytes_planilha_para_cache(conteudo, caminho.name)
            _ARQUIVOS_CARREGADOS.add(chave)
            logger.info("Planilha base carregada automaticamente: %s", caminho.name)
        except Exception as exc:
            logger.warning("Não foi possível carregar planilha %s: %s", caminho, exc)


def buscar_processos_na_planilha_base(cnpj: str) -> List[ProcessoJudicial]:
    _garantir_bases_locais_carregadas()
    cnpj_14 = _limpar_digitos(cnpj)
    lista = _CACHE_PLANILHA.get(cnpj_14, [])
    # Retorna cópias para evitar mutação do cache
    return [p.model_copy(deep=True) for p in lista]
