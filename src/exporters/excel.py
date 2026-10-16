from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models.schemas import EsferaJustica, ProcessoJudicial, ResultadoBuscaCNPJ

# Estilos globais
_FILL_HEADER = PatternFill("solid", fgColor="1F4E79")
_FONT_HEADER = Font(bold=True, color="FFFFFF", size=11)

_FILL_TRABALHISTA = PatternFill("solid", fgColor="DDEEFF")
_FILL_ESTADUAL    = PatternFill("solid", fgColor="D4F0D4")
_FILL_FEDERAL     = PatternFill("solid", fgColor="FFF3CC")
_FILL_OUTROS      = PatternFill("solid", fgColor="F2F2F2")

_ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_ALIGN_LEFT   = Alignment(vertical="center", wrap_text=False)

# (cabeçalho, atributo do ProcessoJudicial, largura)
_COLUNAS_PROCESSO: list[tuple[str, str, int]] = [
    ("Número do Processo",          "numero_formatado",              34),
    ("Tribunal",                    "tribunal",                      12),
    ("Esfera",                      "esfera",                        16),
    ("UF",                          "uf",                             6),
    ("Classe",                      "classe_nome",                   28),
    ("Órgão Julgador",              "orgao_julgador",                36),
    ("Data Ajuizamento",            "data_ajuizamento",              18),
    ("Data Arquivamento",           "data_arquivamento",             18),
    ("Status Detalhado",            "status_detalhado",              28),
    ("Cód. Movimento",              "codigo_movimento_arquivamento", 16),
    ("Último Movimento Relevante",  "ultimo_movimento_relevante",    36),
    ("Polo Empresa",                "polo_empresa",                  14),
    ("Partes Resumo",               "partes_resumo",                 40),
    ("Fonte",                       "fonte_descoberta",              24),
    ("Arquivado",                   "arquivado",                     10),
]


def _fill_esfera(esfera: EsferaJustica) -> PatternFill:
    return {
        EsferaJustica.TRABALHISTA: _FILL_TRABALHISTA,
        EsferaJustica.ESTADUAL:    _FILL_ESTADUAL,
        EsferaJustica.FEDERAL:     _FILL_FEDERAL,
    }.get(esfera, _FILL_OUTROS)


def _fmt(val: object) -> str:
    if val is None:
        return ""
    if isinstance(val, bool):
        return "Sim" if val else "Não"
    if isinstance(val, datetime):
        return val.strftime("%d/%m/%Y")
    # Enum (EsferaJustica, StatusArquivamento) tem .value
    return str(val.value if hasattr(val, "value") else val)


def _escrever_cabecalho(ws, colunas: list[tuple[str, str, int]]) -> None:
    for col_idx, (titulo, _, largura) in enumerate(colunas, start=1):
        cell = ws.cell(row=1, column=col_idx, value=titulo)
        cell.fill = _FILL_HEADER
        cell.font = _FONT_HEADER
        cell.alignment = _ALIGN_CENTER
        ws.column_dimensions[get_column_letter(col_idx)].width = largura
    ws.row_dimensions[1].height = 32


def _escrever_linha_processo(
    ws, row: int, proc: ProcessoJudicial, colunas: list[tuple[str, str, int]]
) -> None:
    fill = _fill_esfera(proc.esfera)
    for col_idx, (_, attr, _) in enumerate(colunas, start=1):
        cell = ws.cell(row=row, column=col_idx, value=_fmt(getattr(proc, attr, None)))
        cell.fill = fill
        cell.alignment = _ALIGN_LEFT


def _criar_aba_processos(
    wb: Workbook, titulo: str, processos: list[ProcessoJudicial]
) -> None:
    ws = wb.create_sheet(title=titulo)
    ws.freeze_panes = "A2"
    _escrever_cabecalho(ws, _COLUNAS_PROCESSO)
    for row_idx, proc in enumerate(processos, start=2):
        _escrever_linha_processo(ws, row_idx, proc, _COLUNAS_PROCESSO)
    if processos:
        ws.auto_filter.ref = ws.dimensions


def _criar_aba_resumo(wb: Workbook, resultado: ResultadoBuscaCNPJ) -> None:
    ws = wb.create_sheet(title="Resumo Executivo", index=0)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 46

    empresa   = resultado.empresa
    contagem  = resultado.contagem_arquivados

    fonte_titulo  = Font(bold=True, size=14, color="1F4E79")
    fonte_secao   = Font(bold=True, size=12, color="1F4E79")
    fonte_rotulo  = Font(bold=True, size=11)
    fonte_valor   = Font(size=11)
    fonte_total   = Font(bold=True, size=11)

    def _linha(r: int, rotulo: str, valor: object, rf: Font = fonte_rotulo, vf: Font = fonte_valor) -> None:
        ws.cell(row=r, column=1, value=rotulo).font = rf
        ws.cell(row=r, column=2, value=_fmt(valor)).font = vf

    # Título
    cell_titulo = ws.cell(row=1, column=1, value="RELATÓRIO DE PROCESSOS ARQUIVADOS")
    cell_titulo.font = fonte_titulo
    cell_titulo.alignment = _ALIGN_CENTER
    ws.merge_cells("A1:B1")
    ws.row_dimensions[1].height = 28

    # Dados da empresa
    ws.cell(row=3, column=1, value="DADOS DA EMPRESA").font = fonte_secao
    ws.merge_cells("A3:B3")

    _linha(4,  "CNPJ:",               empresa.cnpj_formatado)
    _linha(5,  "Razão Social:",        empresa.razao_social)
    _linha(6,  "Nome Fantasia:",       empresa.nome_fantasia)
    _linha(7,  "UF:",                  empresa.uf)
    _linha(8,  "Município:",           empresa.municipio)
    _linha(9,  "Situação Cadastral:",  empresa.situacao_cadastral)
    _linha(10, "Data da Consulta:",    resultado.data_consulta.strftime("%d/%m/%Y %H:%M"))

    # Quadro resumo
    ws.cell(row=12, column=1, value="QUADRO RESUMO – PROCESSOS ARQUIVADOS").font = fonte_secao
    ws.merge_cells("A12:B12")

    for col_idx, label in enumerate(["Esfera", "Quantidade Arquivados"], start=1):
        h = ws.cell(row=13, column=col_idx, value=label)
        h.fill = _FILL_HEADER
        h.font = _FONT_HEADER
        h.alignment = _ALIGN_CENTER

    _linha(14, "Trabalhista", contagem.trabalhista)
    _linha(15, "Estadual",    contagem.estadual)
    _linha(16, "Federal",     contagem.federal)
    _linha(17, "TOTAL ARQUIVADOS",       contagem.total_arquivados,       rf=fonte_total, vf=fonte_total)
    _linha(18, "Total Geral Analisados", contagem.total_geral_analisados, rf=fonte_rotulo, vf=fonte_valor)


def gerar_excel_processos(resultado: ResultadoBuscaCNPJ) -> bytes:
    """
    Gera um arquivo .xlsx em memória com 6 abas:
      1. Resumo Executivo
      2. Processos Arquivados (todos)
      3. Trabalhista (Arquivados)
      4. Estadual (Arquivados)
      5. Federal (Arquivados)
      6. Todos os Processos Analisados
    """
    wb = Workbook()
    wb.remove(wb.active)  # Remove aba padrão "Sheet"

    arquivados = resultado.processos_arquivados

    _criar_aba_resumo(wb, resultado)
    _criar_aba_processos(wb, "Processos Arquivados", arquivados)
    _criar_aba_processos(
        wb, "Trabalhista (Arquivados)",
        [p for p in arquivados if p.esfera == EsferaJustica.TRABALHISTA],
    )
    _criar_aba_processos(
        wb, "Estadual (Arquivados)",
        [p for p in arquivados if p.esfera == EsferaJustica.ESTADUAL],
    )
    _criar_aba_processos(
        wb, "Federal (Arquivados)",
        [p for p in arquivados if p.esfera == EsferaJustica.FEDERAL],
    )
    _criar_aba_processos(wb, "Todos os Processos Analisados", resultado.todos_processos)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
