from __future__ import annotations

from datetime import datetime
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from main import app
from src.logic.archive_classifier import classificar_arquivamento
from src.logic.court_classifier import classificar_esfera_e_tribunal, formatar_numero_cnj
from src.models.schemas import EsferaJustica, MovimentoProcessual, StatusArquivamento

client = TestClient(app)


def test_classificacao_esfera_e_tribunal():
    # Trabalhista TRT-2 (SP) -> J=5, TR=02
    esf, sig, uf = classificar_esfera_e_tribunal("1001234-56.2022.5.02.0001")
    assert esf == EsferaJustica.TRABALHISTA
    assert sig == "TRT2"
    assert uf == "SP"

    # Estadual TJSP -> J=8, TR=26
    esf2, sig2, uf2 = classificar_esfera_e_tribunal("10456781220218260100")
    assert esf2 == EsferaJustica.ESTADUAL
    assert sig2 == "TJSP"
    assert uf2 == "SP"

    # Federal TRF-3 -> J=4, TR=03
    esf3, sig3, uf3 = classificar_esfera_e_tribunal("50098765420204036100")
    assert esf3 == EsferaJustica.FEDERAL
    assert sig3 == "TRF3"
    assert uf3 == "SP"


def test_maquina_estados_arquivamento_cronologico():
    # Caso 1: Arquivado mas os movimentos vieram fora de ordem do DataJud
    movs_fora_de_ordem = [
        MovimentoProcessual(codigo=246, nome="Arquivamento Definitivo", data_hora=datetime(2024, 5, 10, 15, 0)),
        MovimentoProcessual(codigo=26, nome="Distribuição", data_hora=datetime(2021, 1, 10, 10, 0)),
    ]
    arq, status, cod, _, _ = classificar_arquivamento(movs_fora_de_ordem)
    assert arq is True
    assert status == StatusArquivamento.ARQUIVADO_DEFINITIVO
    assert cod == 246

    # Caso 2: Arquivado em 2023 e Desarquivado (893) em 2024 -> NÃO pode contar como arquivado!
    movs_desarquivado = [
        MovimentoProcessual(codigo=26, nome="Distribuição", data_hora=datetime(2022, 1, 10, 10, 0)),
        MovimentoProcessual(codigo=245, nome="Arquivamento Provisório", data_hora=datetime(2023, 6, 1, 12, 0)),
        MovimentoProcessual(codigo=893, nome="Desarquivamento", data_hora=datetime(2024, 2, 15, 14, 0)),
    ]
    arq2, status2, _, _, _ = classificar_arquivamento(movs_desarquivado)
    assert arq2 is False
    assert status2 == StatusArquivamento.DESARQUIVADO


def test_api_consulta_cnpj_e_exportacao_excel():
    resp = client.post(
        "/api/v1/consulta-cnpj",
        json={
            "cnpj": "00.000.000/0001-91",
            "modo_demonstracao": True,
        },
    )
    assert resp.status_code == 200
    dados = resp.json()

    contagem = dados["contagem_arquivados"]
    assert contagem["trabalhista"] == 3
    assert contagem["estadual"] == 3
    assert contagem["federal"] == 2
    assert contagem["total_arquivados"] == 8
    assert contagem["total_geral_analisados"] == 9  # 8 arquivados + 1 desarquivado

    # Testa download do Excel gerado pelo módulo do Claude Code
    job_id = dados["job_id"]
    resp_excel = client.get(f"/api/v1/jobs/{job_id}/excel")
    assert resp_excel.status_code == 200
    assert (
        resp_excel.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

    wb = load_workbook(BytesIO(resp_excel.content))
    assert "Resumo Executivo" in wb.sheetnames
    assert "Processos Arquivados" in wb.sheetnames
    assert "Trabalhista (Arquivados)" in wb.sheetnames
    assert "Estadual (Arquivados)" in wb.sheetnames
    assert "Federal (Arquivados)" in wb.sheetnames
    assert "Todos os Processos Analisados" in wb.sheetnames
